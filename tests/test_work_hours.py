"""Stopda turgan muddat filial ish grafigi bo'yicha hisoblanadi.

Ilgari oddiy ayirma edi: filial yopiq turgan soatlar ham «stopda turdi» deb
sanalardi. Masalan 09:00–03:00 ishlaydigan filialda kechqurun qo'yilgan stop
ertasi kuni yechilsa, oradagi 6 soat yopiqlik ham qo'shilib ketardi.

iiko API grafikni bermaydi (organizations va delivery_restrictions tekshirildi),
shuning uchun u sozlamada saqlanadi va loyiha oynalarida ko'rinmaydi.
"""
from datetime import datetime, timedelta

import pytest

import main
from app import models
from conftest import login


T = lambda *a: datetime(*a)
KUNDUZ = ((9, 0), (3, 0))          # 09:00 -> 03:00 (yarim tundan o'tadi)


def soat(delta):
    return round(delta.total_seconds() / 3600, 2)


def test_grafik_ichida_toliq_sanaladi():
    d = main.working_delta(T(2026, 9, 9, 10, 0), T(2026, 9, 9, 14, 30), *KUNDUZ)
    assert soat(d) == 4.5


def test_yopiq_soatlar_chiqarib_tashlanadi():
    """01:00 da qo'yilib, 11:00 da yechildi. Astronomik 10 soat, lekin
    03:00–09:00 orasida filial yopiq — faqat 4 soat sanaladi."""
    d = main.working_delta(T(2026, 9, 9, 1, 0), T(2026, 9, 9, 11, 0), *KUNDUZ)
    assert soat(d) == 4.0          # 01:00-03:00 (2 s) + 09:00-11:00 (2 s)


def test_toliq_yopiq_oraliq_nol():
    d = main.working_delta(T(2026, 9, 9, 4, 0), T(2026, 9, 9, 8, 0), *KUNDUZ)
    assert soat(d) == 0.0


def test_bir_necha_kun():
    """09.09 20:00 dan 11.09 20:00 gacha — ikki kunlik ish vaqti.
    Har kuni 18 soat ishlaydi (09:00->03:00)."""
    d = main.working_delta(T(2026, 9, 9, 20, 0), T(2026, 9, 11, 20, 0), *KUNDUZ)
    # 09.09 20:00-03:00 = 7s, 10.09 09:00-03:00 = 18s, 11.09 09:00-20:00 = 11s
    assert soat(d) == 36.0


def test_grafik_yoq_bolsa_oddiy_ayirma():
    d = main.working_delta(T(2026, 9, 9, 1, 0), T(2026, 9, 9, 11, 0), None, None)
    assert soat(d) == 10.0


def test_teskari_oraliq_nol():
    d = main.working_delta(T(2026, 9, 9, 11, 0), T(2026, 9, 9, 1, 0), *KUNDUZ)
    assert d == timedelta(0)


def test_hhmm_oqiladi():
    assert main._parse_hhmm("09:00") == (9, 0)
    assert main._parse_hhmm(" 3.05 ") == (3, 5)
    assert main._parse_hhmm("25:00") is None
    assert main._parse_hhmm("") is None
    assert main._parse_hhmm("nima") is None


def test_filial_grafigi_umumiydan_ustun(monkeypatch):
    class B:
        work_from = "10:00"
        work_to = "22:00"
    assert main.branch_work_hours(B()) == ((10, 0), (22, 0))

    class Bosh:
        work_from = ""
        work_to = ""
    monkeypatch.setattr(main, "WORK_HOURS", "08:00-23:00")
    assert main.branch_work_hours(Bosh()) == ((8, 0), (23, 0))

    monkeypatch.setattr(main, "WORK_HOURS", "buzuq")
    assert main.branch_work_hours(Bosh()) == (None, None)


# ---------- Filiallar grafigi spravochnigi («MAXWAY grafik» jadvali) ----------
# Grafik hech qayerda ko'rsatilmaydi va tahrirlanmaydi — u faqat «stopda qancha
# turdi» ni to'g'ri hisoblash uchun kerak.
def test_spravochnik_nom_boyicha():
    assert main.branch_schedule("MW06-NEXT") == "10:00-22:00"
    assert main.branch_schedule("MW12-MAGIC CITY") == "09:00-23:00"
    assert main.branch_schedule("MW01-UNIVERSAM") == "08:00-03:00"


def test_spravochnik_yozilishiga_bogliq_emas():
    """Bo'shliq, tire va katta-kichik harf farq qilmaydi."""
    assert main.branch_schedule("mw03 grand mir") == "08:00-03:00"
    assert main.branch_schedule("MW03-GRANDMIR") == "08:00-03:00"


def test_spravochnik_kod_boyicha():
    """Filial qayta nomlansa ham MW-kodi bo'yicha topiladi."""
    assert main.branch_schedule("MW21-GOLDEN LIFE CENTER") == "10:00-23:00"


def test_spravochnikda_yoq_filial():
    assert main.branch_schedule("Ресторан №12") == ""
    assert main.branch_schedule("") == ""


def test_bazadagi_qiymat_spravochnikdan_ustun():
    """branches.work_from/work_to to'ldirilgan bo'lsa — o'sha ishlaydi."""
    class Oz:
        name, work_from, work_to = "MW06-NEXT", "07:00", "19:00"
    assert main.branch_work_hours(Oz()) == ((7, 0), (19, 0))


def test_spravochnik_umumiy_sozlamadan_ustun(monkeypatch):
    monkeypatch.setattr(main, "WORK_HOURS", "00:00-23:59")

    class Bosh:
        name, work_from, work_to = "MW06-NEXT", "", ""
    assert main.branch_work_hours(Bosh()) == ((10, 0), (22, 0))

    class Notanish:
        name, work_from, work_to = "Ресторан №12", "", ""
    assert main.branch_work_hours(Notanish()) == ((0, 0), (23, 59))


# ---------- «Stopda qancha turdi» ----------
class _Br:
    name, work_from, work_to = "MW06-NEXT", "", ""      # 10:00 — 22:00


class _E:
    def __init__(self, created_at, resolved_at, branch=None):
        self.created_at, self.resolved_at, self.branch = created_at, resolved_at, branch


def test_stop_duration_grafik_boyicha():
    """11.09 20:00 dan 12.09 11:00 gacha. Astronomik 15 soat, lekin filial
    10:00–22:00 ishlaydi: 2 soat (20:00–22:00) + 1 soat (10:00–11:00)."""
    e = _E(T(2026, 9, 11, 20, 0), T(2026, 9, 12, 11, 0), _Br())
    assert main.stop_duration(e) == "3 ч"


def test_stop_duration_yopiq_vaqtda_nol():
    e = _E(T(2026, 9, 12, 2, 0), T(2026, 9, 12, 6, 0), _Br())
    assert main.stop_duration(e) == "меньше минуты"


def test_stop_duration_yechilmagan_bosh():
    assert main.stop_duration(_E(T(2026, 9, 12, 2, 0), None, _Br())) == ""
    assert main.stop_duration(None) == ""


def test_stop_duration_filialsiz_yiqilmaydi(monkeypatch):
    """Filial o'chirilgan bo'lsa ham xato bermaydi — umumiy grafikka tushadi."""
    monkeypatch.setattr(main, "WORK_HOURS", "buzuq")     # grafiksiz = 24/7
    e = _E(T(2026, 9, 12, 2, 0), T(2026, 9, 12, 6, 0), None)
    assert main.stop_duration(e) == "4 ч"


# ---------- Tarix sahifasi, yozuv sahifasi, telegram, eksport ----------
@pytest.fixture
def yechilgan_yozuv(db, seed):
    """11.09 10:00 da stopga tushib, o'sha kuni 14:30 da yechilgan yozuv.
    Filial spravochnikda yo'q — umumiy 09:00–03:00 grafigi ishlaydi, ya'ni
    butun oraliq ish vaqtiga tushadi: 4 soat 30 daqiqa."""
    e = models.StopEntry(
        branch_id=seed["b1"].id, menu_item_id=seed["dishes"][0].id,
        reason="supplier_no_product", source="manual",
        created_at=T(2026, 9, 11, 10, 0), resolved=True,
        resolved_at=T(2026, 9, 11, 14, 30))
    db.add(e)
    db.commit()
    yield e
    db.delete(e)
    db.commit()


def test_tarix_sahifasida_ustun_korinadi(client, seed, yechilgan_yozuv):
    login(client, seed["admin"])
    html = client.get("/stoplist/history?page_size=100").text
    assert "На стопе был" in html
    assert "4 ч 30 мин" in html


def test_yozuv_sahifasida_ham_korinadi(client, seed, yechilgan_yozuv):
    login(client, seed["admin"])
    html = client.get(f"/stoplist/{yechilgan_yozuv.id}").text
    assert "На стопе был" in html
    assert "4 ч 30 мин" in html


def test_telegramda_bitta_taom_muddati(db, seed, yechilgan_yozuv, monkeypatch):
    sent = []
    monkeypatch.setattr(main, "_send_async",
                        lambda ids, text, **kw: sent.append(text))
    main.notify_stop_resolved(db, [yechilgan_yozuv], seed["admin"])
    assert sent and "⏱ На стопе был: 4 ч 30 мин" in sent[0]


def test_telegramda_har_bir_taom_yonida_muddat(db, seed, monkeypatch):
    """Bir nechta taom birga yechilsa — ro'yxatda har birining muddati."""
    sent = []
    monkeypatch.setattr(main, "_send_async",
                        lambda ids, text, **kw: sent.append(text))
    items = []
    for dish, soat in ((seed["dishes"][1], 12), (seed["dishes"][2], 14)):
        e = models.StopEntry(
            branch_id=seed["b1"].id, menu_item_id=dish.id,
            reason="supplier_no_product", source="manual",
            created_at=T(2026, 9, 11, 10, 0), resolved=True,
            resolved_at=T(2026, 9, 11, soat, 0))
        db.add(e)
        items.append(e)
    db.commit()
    try:
        main.notify_stop_resolved(db, items, seed["admin"])
        assert sent
        assert "Картофель фри — 2 ч" in sent[0]
        assert "Кола 0.5 — 4 ч" in sent[0]
    finally:
        for e in items:
            db.delete(e)
        db.commit()


def test_eksportda_ham_ustun_bor(client, seed, yechilgan_yozuv):
    """Tarix eksporti sahifadagi jadvalning aynan o'zi bo'lishi kerak."""
    import io as _io
    import openpyxl
    login(client, seed["admin"])
    r = client.get("/stoplist/export?mode=history")
    assert r.status_code == 200
    ws = openpyxl.load_workbook(_io.BytesIO(r.content)).worksheets[0]
    headers = [c.value for c in ws[1]]
    assert headers[-2:] == ["Убрано", "На стопе был"]
    vals = {row[2].value: row[-1].value for row in ws.iter_rows(min_row=2)}
    assert vals.get("Бургер Классик") == "4 ч 30 мин"


def test_aktiv_royxat_eksportida_ustun_yoq(client, seed):
    """Hali yechilmagan yozuvda muddat ma'nosiz — ustun faqat tarixda."""
    import io as _io
    import openpyxl
    login(client, seed["admin"])
    r = client.get("/stoplist/export?mode=active")
    ws = openpyxl.load_workbook(_io.BytesIO(r.content)).worksheets[0]
    assert "На стопе был" not in [c.value for c in ws[1]]


def test_grafik_admin_panelida_korinmaydi(client, seed):
    """Grafik — ichki hisob vositasi: hech bir oynada ko'rsatilmaydi va
    tahrirlanmaydi (faqat «На стопе был» soni ko'rinadi)."""
    login(client, seed["admin"])
    html = client.get("/admin").text
    assert "work_hours" not in html
    assert "График работы" not in html
